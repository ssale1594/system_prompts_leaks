# -*- coding: utf-8 -*-
"""Improve Only_Recieving.xlsx: styling, freeze panes, filters, column groups,
conditional formatting, expiry highlighting, dashboard + instructions sheets."""
import re
from datetime import date, datetime
import calendar

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter

SRC = 'original.xlsx'
OUT = 'Only_Receiving_Improved.xlsx'
TODAY = date(2026, 7, 6)
NEAR_DAYS = 180

wb = openpyxl.load_workbook(SRC)
wbv = openpyxl.load_workbook(SRC, data_only=True)  # cached values

# ---------------------------------------------------------------- helpers
MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_name) if m}

def parse_date(v):
    """Best-effort parse of the many date formats used in the file."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    if not s or s.upper() in ('NA', 'N/A', '-'):
        return None
    # dd.mm.yyyy / dd/mm/yyyy / dd-mm-yyyy
    m = re.fullmatch(r'(\d{1,2})[./-](\d{1,2})[./-](\d{4})', s)
    if m:
        d_, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return date(y, mo, d_)
        except ValueError:
            return None
    # mm/yyyy or mm.yyyy or mm-yyyy -> end of month
    m = re.fullmatch(r'(\d{1,2})[./-](\d{4})', s)
    if m:
        mo, y = int(m.group(1)), int(m.group(2))
        if 1 <= mo <= 12:
            return date(y, mo, calendar.monthrange(y, mo)[1])
        return None
    # "February.-2026" / "september.-2027" / "Feb-2026"
    m = re.fullmatch(r'([A-Za-z]+)\W*[-. ]\W*(\d{4})', s)
    if m:
        name = m.group(1).lower()
        y = int(m.group(2))
        mo = MONTHS.get(name)
        if mo is None:
            for full, idx in MONTHS.items():
                if full.startswith(name[:3]):
                    mo = idx
                    break
        if mo:
            return date(y, mo, calendar.monthrange(y, mo)[1])
    return None

def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None

# ---------------------------------------------------------------- clean-up
# 1. Trim the thousands of junk columns in Off -Records (real data ends at BK=63)
off = wb['Off -Records']
if off.max_column > 63:
    off.delete_cols(64, off.max_column - 63)

# 2. Add the missing header in RM & PM
wb['RM & PM']['A1'] = 'Item Category'

# 3. Fix sheet-name typo / trailing space (no cross-sheet formulas exist)
wb['Hisorical I & A'].title = 'Historical I & A'
wb['Inner & Assembled '].title = 'Inner & Assembled'
wbv_names = {'Hisorical I & A': 'Historical I & A', 'Inner & Assembled ': 'Inner & Assembled'}

# ---------------------------------------------------------------- layout config
# sections: list of (first_col, last_col, fill) for header colouring
BLUE, TEAL, GREEN, ORANGE, GRAY = '305496', '1F7A6D', '548235', 'C65911', '7F7F7F'

CFG = {
    'RM & PM':          dict(id=(1, 10), dates=(11, 17), qty=(18, 25), status=(26, 31),
                             txn=(32, 67), released=27, quarantine=28, rejected=31,
                             code=5, name=6, lot=7, exp=13, lastcol=70),
    'Historical R&PM':  dict(id=(1, 10), dates=(11, 17), qty=(18, 25), status=(26, 31),
                             txn=(32, 67), released=27, quarantine=28, rejected=31,
                             code=5, name=6, lot=7, exp=13, lastcol=70),
    'Off -Records':     dict(id=(1, 10), dates=(11, 13), qty=(14, 21), status=(22, 27),
                             txn=(28, 63), released=23, quarantine=24, rejected=27,
                             code=5, name=6, lot=7, exp=12, lastcol=63),
    'Stagnant I & A':   dict(id=(1, 9), dates=(7, 8), qty=(10, 13), status=(14, 19),
                             txn=(20, 55), released=15, quarantine=16, rejected=19,
                             code=4, name=6, lot=5, exp=8, lastcol=55),
    'Historical I & A': dict(id=(1, 9), dates=(7, 8), qty=(10, 13), status=(14, 19),
                             txn=(20, 55), released=15, quarantine=16, rejected=19,
                             code=4, name=6, lot=5, exp=8, lastcol=55),
    'Inner & Assembled': dict(id=(1, 9), dates=(7, 8), qty=(10, 13), status=(14, 19),
                              txn=(20, 55), released=15, quarantine=16, rejected=19,
                              code=4, name=6, lot=5, exp=8, lastcol=55),
}

TAB_COLORS = {
    'RM & PM': '217346', 'Inner & Assembled': '217346',
    'Off -Records': '7030A0', 'Stagnant I & A': 'ED7D31',
    'Historical R&PM': 'A6A6A6', 'Historical I & A': 'A6A6A6',
}

CURRENT_SHEETS = ['RM & PM', 'Off -Records', 'Inner & Assembled', 'Stagnant I & A']

thin = Side(style='thin', color='D9D9D9')
hdr_font = Font(bold=True, color='FFFFFF', size=10)
hdr_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

RED_FILL = PatternFill('solid', fgColor='FFC7CE')
RED_FONT = Font(color='9C0006')
YEL_FILL = PatternFill('solid', fgColor='FFEB9C')
YEL_FONT = Font(color='9C6500')
ORA_FILL = PatternFill('solid', fgColor='FCE4D6')
ZEBRA = PatternFill('solid', fgColor='F2F6FA')

expired_rows = []   # (sheet, code, name, lot/batch, exp_date, released, status)
neg_released = []

for name, cfg in CFG.items():
    ws = wb[name]
    lastcol = cfg['lastcol']
    lastrow = ws.max_row
    LC = get_column_letter(lastcol)

    # ---- header styling by section
    def fill_range(c1, c2, color):
        f = PatternFill('solid', fgColor=color)
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = f
            cell.font = hdr_font
            cell.alignment = hdr_align
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill_range(*cfg['id'], BLUE)
    fill_range(*cfg['dates'], TEAL)
    fill_range(*cfg['qty'], GREEN)
    fill_range(*cfg['status'], ORANGE)
    fill_range(*cfg['txn'], GRAY)
    ws.row_dimensions[1].height = 34

    # ---- number the 12 repeated transaction triplets
    t0, t1 = cfg['txn']
    n = 0
    for c in range(t0, t1 + 1, 3):
        n += 1
        ws.cell(row=1, column=c).value = f'Date ({n})'
        ws.cell(row=1, column=c + 1).value = f'Doc No. ({n})'
        ws.cell(row=1, column=c + 2).value = f'Qty ({n})'

    # ---- freeze panes + autofilter
    ws.freeze_panes = 'G2'
    ws.auto_filter.ref = f'A1:{LC}{lastrow}'

    # ---- collapse transaction columns into an outline group
    for c in range(t0, t1 + 1):
        cd = ws.column_dimensions[get_column_letter(c)]
        cd.outline_level = 1
        cd.hidden = True
    ws.column_dimensions[get_column_letter(t1 + 1)].collapsed = True
    ws.sheet_properties.outlinePr.summaryRight = True

    # ---- column widths
    for c in range(1, lastcol + 1):
        ws.column_dimensions[get_column_letter(c)].width = 12
    ws.column_dimensions[get_column_letter(cfg['name'])].width = 42
    if name in ('RM & PM', 'Historical R&PM', 'Off -Records'):
        ws.column_dimensions['H'].width = 26   # manufacturer
        ws.column_dimensions['I'].width = 22   # supplier
    for c in range(t0, t1 + 1):
        cd = ws.column_dimensions[get_column_letter(c)]
        cd.outline_level = 1
        cd.hidden = True

    # ---- conditional formatting
    data_rng = f'A2:{LC}{lastrow}'
    ws.conditional_formatting.add(
        data_rng, FormulaRule(formula=['MOD(ROW(),2)=0'], fill=ZEBRA, stopIfTrue=False))
    rel = get_column_letter(cfg['released'])
    qua = get_column_letter(cfg['quarantine'])
    rej = get_column_letter(cfg['rejected'])
    ws.conditional_formatting.add(
        f'{rel}2:{rel}{lastrow}',
        CellIsRule(operator='lessThan', formula=['0'], fill=RED_FILL, font=RED_FONT))
    ws.conditional_formatting.add(
        f'{qua}2:{qua}{lastrow}',
        CellIsRule(operator='greaterThan', formula=['0'], fill=YEL_FILL, font=YEL_FONT))
    ws.conditional_formatting.add(
        f'{rej}2:{rej}{lastrow}',
        CellIsRule(operator='greaterThan', formula=['0'], fill=RED_FILL, font=RED_FONT))

    # ---- tab colour
    ws.sheet_properties.tabColor = TAB_COLORS[name]

    # ---- static expiry scan (current sheets only) using cached values
    if name in CURRENT_SHEETS:
        vname = name if name in wbv.sheetnames else \
            [k for k, v in wbv_names.items() if v == name][0]
        wsv = wbv[vname]
        for r in range(2, lastrow + 1):
            code = ws.cell(row=r, column=cfg['code']).value
            if code is None or not str(code).strip():
                continue
            expd = parse_date(ws.cell(row=r, column=cfg['exp']).value)
            relv = num(wsv.cell(row=r, column=cfg['released']).value)
            if relv is not None and relv < -0.001:
                relv = round(relv, 3)
                neg_released.append((name, str(code).strip(),
                                     str(ws.cell(row=r, column=cfg['name']).value or '').strip(),
                                     relv, r))
            if expd is None or relv is None or relv <= 0.001:
                continue
            days = (expd - TODAY).days
            if days < 0:
                status = 'منتهي الصلاحية'
            elif days <= NEAR_DAYS:
                status = f'ينتهي خلال {days} يوم'
            else:
                continue
            lot = ws.cell(row=r, column=cfg['lot']).value
            expired_rows.append((name, str(code).strip(),
                                 str(ws.cell(row=r, column=cfg['name']).value or '').strip()[:60],
                                 str(lot or '').strip(), expd, round(relv, 3), status, days))
            # highlight the expiry cell in the data sheet (snapshot)
            cell = ws.cell(row=r, column=cfg['exp'])
            cell.fill = RED_FILL if days < 0 else ORA_FILL
            if days < 0:
                cell.font = RED_FONT

expired_rows.sort(key=lambda x: x[7])

# ---------------------------------------------------------------- dashboard
dash = wb.create_sheet('الملخص', 0)
dash.sheet_view.rightToLeft = True
dash.sheet_properties.tabColor = '1F4E79'
dash.sheet_view.showGridLines = False

title_font = Font(bold=True, size=16, color='1F4E79')
sec_font = Font(bold=True, size=12, color='FFFFFF')
sec_fill = PatternFill('solid', fgColor='305496')
lbl_font = Font(bold=True, size=10)
note_font = Font(size=9, color='808080', italic=True)
box_border = Border(left=thin, right=thin, top=thin, bottom=thin)

dash['B2'] = 'لوحة ملخص المخزون والاستلام'
dash['B2'].font = title_font
dash['B3'] = f'آخر تحليل للصلاحية: {TODAY.strftime("%d.%m.%Y")} — الإحصائيات في الجدول الأول تتحدّث تلقائياً مع تعديل البيانات'
dash['B3'].font = note_font

# --- live per-sheet stats
row = 5
dash.cell(row=row, column=2, value='إحصائيات الأوراق (تلقائية التحديث)').font = sec_font
dash.cell(row=row, column=2).fill = sec_fill
for c in range(3, 7):
    dash.cell(row=row, column=c).fill = sec_fill
row += 1
headers = ['الورقة', 'عدد الدفعات', 'دفعات برصيد حالي', 'دفعات برصيد سالب (خطأ)', 'ملاحظة']
for i, h in enumerate(headers):
    cell = dash.cell(row=row, column=2 + i, value=h)
    cell.font = Font(bold=True, color='FFFFFF', size=10)
    cell.fill = PatternFill('solid', fgColor='7F7F7F')
    cell.alignment = Alignment(horizontal='center', wrap_text=True)
    cell.border = box_border

sheet_meta = [
    ('RM & PM', 'E', 'AA', 'المواد الخام ومواد التعبئة الحالية'),
    ('Off -Records', 'E', 'W', 'أصناف خارج السجلات الرسمية'),
    ('Inner & Assembled', 'D', 'O', 'المنتجات الداخلية والمجمّعة الحالية'),
    ('Stagnant I & A', 'D', 'O', 'أصناف راكدة'),
    ('Historical R&PM', 'E', 'AA', 'أرشيف المواد الخام والتعبئة'),
    ('Historical I & A', 'D', 'O', 'أرشيف المنتجات الداخلية والمجمّعة'),
]
for sname, codecol, relcol, desc in sheet_meta:
    row += 1
    lr = wb[sname].max_row
    q = f"'{sname}'"
    vals = [sname,
            f'=COUNTA({q}!{codecol}2:{codecol}{lr})',
            f'=COUNTIF({q}!{relcol}2:{relcol}{lr},">0")',
            f'=COUNTIF({q}!{relcol}2:{relcol}{lr},"<-0.001")',
            desc]
    for i, v in enumerate(vals):
        cell = dash.cell(row=row, column=2 + i, value=v)
        cell.border = box_border
        if i in (1, 2, 3):
            cell.alignment = Alignment(horizontal='center')

# --- expiry alerts (static snapshot)
row += 2
dash.cell(row=row, column=2,
          value=f'تنبيهات الصلاحية — أصناف برصيد حالي منتهية أو تنتهي خلال {NEAR_DAYS} يوم '
                f'(لقطة بتاريخ {TODAY.strftime("%d.%m.%Y")})').font = sec_font
dash.cell(row=row, column=2).fill = PatternFill('solid', fgColor='C65911')
for c in range(3, 9):
    dash.cell(row=row, column=c).fill = PatternFill('solid', fgColor='C65911')
row += 1
exp_headers = ['الورقة', 'الكود', 'اسم الصنف', 'رقم التشغيلة/الكنترول', 'تاريخ الانتهاء',
               'الرصيد الحالي', 'الحالة']
for i, h in enumerate(exp_headers):
    cell = dash.cell(row=row, column=2 + i, value=h)
    cell.font = Font(bold=True, color='FFFFFF', size=10)
    cell.fill = PatternFill('solid', fgColor='7F7F7F')
    cell.alignment = Alignment(horizontal='center', wrap_text=True)
    cell.border = box_border

if not expired_rows:
    row += 1
    dash.cell(row=row, column=2, value='لا توجد أصناف منتهية أو قريبة الانتهاء برصيد حالي ✓')
else:
    for sname, code, iname, lot, expd, relv, status, days in expired_rows:
        row += 1
        vals = [sname, code, iname, lot, expd.strftime('%d.%m.%Y'), relv, status]
        for i, v in enumerate(vals):
            cell = dash.cell(row=row, column=2 + i, value=v)
            cell.border = box_border
            if i in (4, 5, 6):
                cell.alignment = Alignment(horizontal='center')
        fill = RED_FILL if days < 0 else ORA_FILL
        for i in range(len(vals)):
            dash.cell(row=row, column=2 + i).fill = fill
        if days < 0:
            dash.cell(row=row, column=8).font = RED_FONT

# --- negative balances (data-entry errors)
if neg_released:
    row += 2
    dash.cell(row=row, column=2,
              value='أرصدة سالبة تحتاج مراجعة (غالباً خطأ إدخال)').font = sec_font
    dash.cell(row=row, column=2).fill = PatternFill('solid', fgColor='9C0006')
    for c in range(3, 7):
        dash.cell(row=row, column=c).fill = PatternFill('solid', fgColor='9C0006')
    row += 1
    for i, h in enumerate(['الورقة', 'الكود', 'اسم الصنف', 'الرصيد', 'رقم الصف']):
        cell = dash.cell(row=row, column=2 + i, value=h)
        cell.font = Font(bold=True, color='FFFFFF', size=10)
        cell.fill = PatternFill('solid', fgColor='7F7F7F')
        cell.border = box_border
    for sname, code, iname, relv, r in neg_released:
        row += 1
        for i, v in enumerate([sname, code, iname[:60], relv, r]):
            cell = dash.cell(row=row, column=2 + i, value=v)
            cell.border = box_border
            cell.fill = RED_FILL

widths = {'A': 2, 'B': 20, 'C': 14, 'D': 50, 'E': 22, 'F': 14, 'G': 14, 'H': 20, 'I': 12}
for col, w in widths.items():
    dash.column_dimensions[col].width = w

# ---------------------------------------------------------------- instructions
ins = wb.create_sheet('دليل الاستخدام', 1)
ins.sheet_view.rightToLeft = True
ins.sheet_properties.tabColor = '1F4E79'
ins.sheet_view.showGridLines = False
ins.column_dimensions['A'].width = 2
ins.column_dimensions['B'].width = 110

lines = [
    ('دليل استخدام ملف الاستلام والمخزون', 'title'),
    ('', None),
    ('ما الجديد في هذه النسخة؟', 'sec'),
    ('1. لوحة "الملخص": إحصائيات لكل ورقة + تنبيهات الأصناف المنتهية أو القريبة من انتهاء الصلاحية + الأرصدة السالبة.', None),
    ('2. أعمدة حركات الصرف (12 حركة × تاريخ/مستند/كمية) مجمّعة ومطوية افتراضياً — اضغط زر (+) أعلى يمين الأعمدة لعرضها، و(−) لإخفائها.', None),
    ('3. تجميد الصفوف والأعمدة: صف العناوين وأعمدة التعريف (حتى اسم الصنف) تبقى ظاهرة أثناء التمرير.', None),
    ('4. فلاتر تلقائية على كل الأعمدة: اضغط السهم في رأس أي عمود للتصفية أو الترتيب.', None),
    ('5. ترقيم حركات الصرف: Date (1), Doc No. (1), Qty (1) ... بدل التكرار بدون تمييز.', None),
    ('6. حذف آلاف الأعمدة الفارغة من ورقة Off -Records (كانت تبطئ الملف).', None),
    ('', None),
    ('دلالة ألوان رؤوس الأعمدة', 'sec'),
    ('أزرق: بيانات تعريف الصنف (الكود، الاسم، المورد...) | أخضر مزرق: التواريخ | أخضر: الكميات المستلمة | برتقالي: أرصدة الحالة (المفرج عنه، الحجر، المرفوض) | رمادي: حركات الصرف.', None),
    ('', None),
    ('التنسيق الشرطي (يتحدّث تلقائياً)', 'sec'),
    ('• الرصيد الحالي أحمر = رصيد سالب (خطأ إدخال يحتاج مراجعة).', None),
    ('• عمود الحجر (Quarantine) أصفر = توجد كمية محجوزة.', None),
    ('• عمود المرفوض (Rejected) أحمر = توجد كمية مرفوضة.', None),
    ('', None),
    ('تظليل تواريخ الانتهاء (لقطة ثابتة)', 'sec'),
    (f'خلايا تاريخ الانتهاء المظللة بالأحمر = منتهية، وبالبرتقالي = تنتهي خلال {NEAR_DAYS} يوم، وذلك حسب تحليل بتاريخ {TODAY.strftime("%d.%m.%Y")} وللأصناف التي لها رصيد حالٍ فقط. هذا التظليل لا يتحدّث تلقائياً لأن التواريخ مكتوبة كنص بصيغ مختلفة.', None),
    ('', None),
    ('توصية مهمة لتسهيل الملف مستقبلاً', 'sec'),
    ('وحّدوا كتابة التواريخ بصيغة واحدة (يفضَّل تاريخ إكسل حقيقي مثل 06/07/2026) بدل الخلط بين 11.06.2023 و 11/06/2023 و February-2023 — عندها يمكن جعل تنبيهات الصلاحية تلقائية بالكامل.', None),
    ('', None),
    ('وصف الأوراق', 'sec'),
    ('RM & PM: استلامات المواد الخام ومواد التعبئة الحالية | Off -Records: أصناف خارج السجلات | Inner & Assembled: المنتجات الداخلية والمجمّعة | Stagnant I & A: الأصناف الراكدة | Historical R&PM و Historical I & A: الأرشيف.', None),
]
r = 2
for text, kind in lines:
    cell = ins.cell(row=r, column=2, value=text)
    if kind == 'title':
        cell.font = Font(bold=True, size=16, color='1F4E79')
    elif kind == 'sec':
        cell.font = Font(bold=True, size=12, color='C65911')
    else:
        cell.font = Font(size=11)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        if text:
            ins.row_dimensions[r].height = max(18, 16 * (len(text) // 95 + 1))
    r += 1

# ---------------------------------------------------------------- sheet order
order = ['الملخص', 'دليل الاستخدام', 'RM & PM', 'Off -Records', 'Inner & Assembled',
         'Stagnant I & A', 'Historical R&PM', 'Historical I & A']
wb._sheets = [wb[n] for n in order]
wb.active = 0

wb.save(OUT)
print('saved', OUT)
print('expiry alerts:', len(expired_rows), '| negative balances:', len(neg_released))
for e in expired_rows[:15]:
    print('  ', e[0], e[1], e[4], e[5], e[6])
